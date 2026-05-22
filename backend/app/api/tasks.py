import io
import uuid
from typing import Optional, List
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, status
from fastapi.responses import Response
from sqlalchemy.orm import Session
from sqlalchemy import or_
from openpyxl import Workbook, load_workbook

from app.core.deps import (
    get_db, get_current_user, require_task_access, apply_task_scope,
    can_delete_task, visible_department_ids,
)
from app.models.task import Task, TaskComment
from app.models.department import Department
from app.models.user import User
from app.models.enums import TaskType, TaskPriority, TaskStatus, UserRole
from app.schemas.task import TaskCreate, TaskUpdate, CommentCreate, CommentUpdate, TaskOut, TaskDetailOut
from app.services.task_service import (
    compute_period_bucket, infer_child_type, log_history,
    recompute_parent_progress, get_parent_chain, build_tree,
)
from app.services.notifications import notify_task_assignee, task_link

# ── RU ↔ Enum maps for Excel import / export ────────────────────────────

_TYPE_RU: dict[TaskType, str] = {
    TaskType.GOAL: "Цель",
    TaskType.EPIC: "Эпик",
    TaskType.TASK: "Задача",
    TaskType.SUBTASK: "Подзадача",
}
_TYPE_FROM_LABEL: dict[str, TaskType] = {
    v.lower(): k for k, v in _TYPE_RU.items()
}
_TYPE_FROM_LABEL.update({e.value.lower(): e for e in TaskType})

_STATUS_RU: dict[TaskStatus, str] = {
    TaskStatus.NEW: "Новая",
    TaskStatus.IN_PROGRESS: "В работе",
    TaskStatus.REVIEW: "Ревью",
    TaskStatus.DONE: "Готово",
    TaskStatus.OVERDUE: "Просрочена",
}
_STATUS_FROM_LABEL: dict[str, TaskStatus] = {
    v.lower(): k for k, v in _STATUS_RU.items()
}
_STATUS_FROM_LABEL.update({e.value.lower(): e for e in TaskStatus})

_PRIORITY_RU: dict[TaskPriority, str] = {
    TaskPriority.LOW: "Низкий",
    TaskPriority.MEDIUM: "Средний",
    TaskPriority.HIGH: "Высокий",
    TaskPriority.CRITICAL: "Критический",
}
_PRIORITY_FROM_LABEL: dict[str, TaskPriority] = {
    v.lower(): k for k, v in _PRIORITY_RU.items()
}
_PRIORITY_FROM_LABEL.update({e.value.lower(): e for e in TaskPriority})

# Header synonyms (lower-cased) → canonical key
_HEADER_SYNONYMS: dict[str, str] = {
    "тип": "type", "type": "type",
    "название": "title", "title": "title",
    "описание": "description", "description": "description",
    "статус": "status", "status": "status",
    "приоритет": "priority", "priority": "priority",
    "отдел": "department", "department": "department",
    "исполнитель": "assignee", "assignee": "assignee",
    "дата начала": "start_date", "start_date": "start_date",
    "срок": "due_date", "due_date": "due_date",
    "прогресс": "progress", "progress": "progress",
}

router = APIRouter(prefix="/tasks", tags=["tasks"])


def _task_out(t: Task) -> dict:
    children_count = len(t.children) if t.children else 0
    return {
        "id": t.id,
        "type": t.type,
        "title": t.title,
        "description": t.description,
        "parent_id": t.parent_id,
        "start_date": t.start_date,
        "due_date": t.due_date,
        "priority": t.priority,
        "status": t.status,
        "progress": t.progress,
        "assigned_department_id": t.assigned_department_id,
        "assigned_department_name": t.assigned_department.name if t.assigned_department else None,
        "assignee_id": t.assignee_id,
        "assignee_name": t.assignee.full_name if t.assignee else None,
        "created_by_id": t.created_by_id,
        "created_by_name": t.created_by.full_name if t.created_by else None,
        "period_bucket": compute_period_bucket(t.start_date, t.due_date),
        "children_count": children_count,
        "created_at": t.created_at,
        "updated_at": t.updated_at,
    }


###############################################################################
# Excel export / import
###############################################################################

_EXPORT_HEADERS = [
    "Тип", "Название", "Описание", "Статус", "Приоритет",
    "Отдел", "Исполнитель", "Дата начала", "Срок", "Прогресс",
]


@router.get("/export")
def export_tasks(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    tasks = apply_task_scope(db.query(Task), current_user, db).order_by(Task.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Задачи"
    ws.append(_EXPORT_HEADERS)

    for t in tasks:
        ws.append([
            _TYPE_RU.get(t.type, t.type.value if t.type else ""),
            t.title,
            t.description or "",
            _STATUS_RU.get(t.status, t.status.value if t.status else ""),
            _PRIORITY_RU.get(t.priority, t.priority.value if t.priority else ""),
            t.assigned_department.name if t.assigned_department else "",
            t.assignee.full_name if t.assignee else "",
            t.start_date.isoformat() if t.start_date else "",
            t.due_date.isoformat() if t.due_date else "",
            t.progress,
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="tasks.xlsx"'},
    )


def _parse_date(value) -> Optional[date]:
    """Accept a date/datetime cell or an ISO-ish string."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


@router.post("/import")
def import_tasks(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    # RBAC: EMPLOYEE forbidden
    if current_user.role == UserRole.EMPLOYEE:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Сотрудник не может импортировать задачи",
        )

    # Pre-load lookup caches
    dept_by_name: dict[str, Department] = {
        d.name.lower(): d for d in db.query(Department).all()
    }
    user_by_name: dict[str, User] = {
        u.full_name.lower(): u for u in db.query(User).all()
    }
    user_by_email: dict[str, User] = {
        u.email.lower(): u for u in db.query(User).all()
    }

    allowed_depts = visible_department_ids(db, current_user)  # None for ADMIN

    try:
        wb = load_workbook(filename=io.BytesIO(file.file.read()), data_only=True)
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Не удалось прочитать файл как .xlsx",
        )

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "skipped": 0, "errors": []}

    # Build header map (row 0)
    raw_headers = rows[0]
    col_map: dict[str, int] = {}
    for idx, h in enumerate(raw_headers):
        if h is None:
            continue
        key = _HEADER_SYNONYMS.get(str(h).strip().lower())
        if key and key not in col_map:
            col_map[key] = idx

    def cell(row_data, key: str):
        idx = col_map.get(key)
        if idx is None or idx >= len(row_data):
            return None
        v = row_data[idx]
        if isinstance(v, str):
            v = v.strip()
        return v if v != "" else None

    created = 0
    skipped = 0
    errors: list[dict] = []

    for row_num, row_data in enumerate(rows[1:], start=2):
        title = cell(row_data, "title")
        if not title:
            skipped += 1
            errors.append({"row": row_num, "message": "Пропущено: отсутствует название"})
            continue

        # Type
        raw_type = cell(row_data, "type")
        task_type = _TYPE_FROM_LABEL.get(str(raw_type).lower()) if raw_type else TaskType.TASK
        if task_type is None:
            task_type = TaskType.TASK

        # Status
        raw_status = cell(row_data, "status")
        task_status = _STATUS_FROM_LABEL.get(str(raw_status).lower()) if raw_status else TaskStatus.NEW
        if task_status is None:
            task_status = TaskStatus.NEW

        # Priority
        raw_priority = cell(row_data, "priority")
        task_priority = _PRIORITY_FROM_LABEL.get(str(raw_priority).lower()) if raw_priority else TaskPriority.MEDIUM
        if task_priority is None:
            task_priority = TaskPriority.MEDIUM

        # Department
        raw_dept = cell(row_data, "department")
        dept: Optional[Department] = None
        if raw_dept:
            dept = dept_by_name.get(str(raw_dept).lower())
            if dept is None:
                skipped += 1
                errors.append({"row": row_num, "message": f"Отдел не найден: {raw_dept}"})
                continue

        dept_id = dept.id if dept else None

        # LEAD RBAC: department must be within visible subtree
        if current_user.role == UserRole.LEAD:
            if dept_id is not None and allowed_depts is not None and dept_id not in allowed_depts:
                skipped += 1
                errors.append({"row": row_num, "message": f"Нет доступа к отделу: {raw_dept}"})
                continue
            if dept_id is None:
                dept_id = current_user.department_id

        # Assignee
        raw_assignee = cell(row_data, "assignee")
        assignee: Optional[User] = None
        if raw_assignee:
            key = str(raw_assignee).lower()
            assignee = user_by_name.get(key) or user_by_email.get(key)

        # Dates
        start_dt = _parse_date(cell(row_data, "start_date"))
        due_dt = _parse_date(cell(row_data, "due_date"))

        # Progress
        raw_progress = cell(row_data, "progress")
        progress = 0
        if raw_progress is not None:
            try:
                progress = max(0, min(100, int(float(str(raw_progress)))))
            except (ValueError, TypeError):
                progress = 0

        task = Task(
            id=uuid.uuid4(),
            type=task_type,
            title=str(title),
            description=cell(row_data, "description") or None,
            start_date=start_dt,
            due_date=due_dt,
            priority=task_priority,
            status=task_status,
            progress=progress,
            assigned_department_id=dept_id,
            assignee_id=assignee.id if assignee else None,
            created_by_id=current_user.id,
        )
        db.add(task)
        db.flush()

        log_history(db, task.id, "created", current_user.id, {"title": task.title, "type": task.type.value})
        created += 1

    db.commit()
    return {"created": created, "skipped": skipped, "errors": errors}


@router.get("/tree")
def get_tree(
    root_id: Optional[uuid.UUID] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return build_tree(db, root_id, current_user=current_user)


@router.get("")
def list_tasks(
    type_filter: Optional[str] = Query(None, alias="type"),
    department_id: Optional[uuid.UUID] = None,
    assignee_id: Optional[uuid.UUID] = None,
    created_by_id: Optional[uuid.UUID] = None,
    status_filter: Optional[str] = Query(None, alias="status"),
    priority_filter: Optional[str] = Query(None, alias="priority"),
    parent_id: Optional[str] = None,
    q: Optional[str] = None,
    due_before: Optional[date] = None,
    due_after: Optional[date] = None,
    limit: int = Query(200, le=1000),
    offset: int = 0,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = apply_task_scope(db.query(Task), current_user, db)

    # Type filter is a UNION (OR) over the selected types — selecting
    # "Цель" + "Эпик" returns tasks that are GOAL *or* EPIC, matching how
    # status/priority already behave. (Previously a single-value param,
    # which made multi-select either break or behave like an intersection.)
    if type_filter:
        types = [t.strip() for t in type_filter.split(",") if t.strip()]
        if types:
            query = query.filter(Task.type.in_(types))
    if department_id:
        query = query.filter(Task.assigned_department_id == department_id)
    if assignee_id:
        query = query.filter(Task.assignee_id == assignee_id)
    if created_by_id:
        query = query.filter(Task.created_by_id == created_by_id)
    if status_filter:
        statuses = [s.strip() for s in status_filter.split(",") if s.strip()]
        query = query.filter(Task.status.in_(statuses))
    if priority_filter:
        priorities = [p.strip() for p in priority_filter.split(",") if p.strip()]
        query = query.filter(Task.priority.in_(priorities))
    if parent_id is not None:
        if parent_id == "null":
            query = query.filter(Task.parent_id.is_(None))
        else:
            try:
                parent_uuid = uuid.UUID(parent_id)
            except ValueError:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid parent_id")
            query = query.filter(Task.parent_id == parent_uuid)
    if q:
        pattern = f"%{q}%"
        query = query.filter(or_(Task.title.ilike(pattern), Task.description.ilike(pattern)))
    if due_before:
        query = query.filter(Task.due_date <= due_before)
    if due_after:
        query = query.filter(Task.due_date >= due_after)

    query = query.order_by(Task.created_at.desc())
    tasks = query.offset(offset).limit(limit).all()
    return [_task_out(t) for t in tasks]


@router.post("", status_code=status.HTTP_201_CREATED)
def create_task(
    body: TaskCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    # ── RBAC: only ADMIN and LEAD may create tasks/subtasks ──────────────
    if current_user.role == UserRole.EMPLOYEE:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Сотрудник не может создавать задачи",
        )

    task_type = body.type or TaskType.TASK
    if body.parent_id:
        parent = require_task_access(body.parent_id, db, current_user)
        task_type = infer_child_type(parent.type)

    # Auto-set department from assignee if not provided
    dept_id = body.assigned_department_id
    if not dept_id and body.assignee_id:
        assignee = db.get(User, body.assignee_id)
        if assignee and assignee.department_id:
            dept_id = assignee.department_id

    # LEAD may only create within their own department subtree and assign
    # people from it. ADMIN is unrestricted (visible_department_ids → None).
    if current_user.role == UserRole.LEAD:
        allowed_depts = visible_department_ids(db, current_user)
        if dept_id is None:
            dept_id = current_user.department_id
        if allowed_depts is not None and dept_id is not None and dept_id not in allowed_depts:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Руководитель может создавать задачи только в своём отделе",
            )
        if body.assignee_id:
            a = db.get(User, body.assignee_id)
            if a and a.department_id and allowed_depts is not None and a.department_id not in allowed_depts:
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN,
                    detail="Можно назначать только сотрудников своего отдела",
                )

    task = Task(
        id=uuid.uuid4(),
        type=task_type,
        title=body.title,
        description=body.description,
        parent_id=body.parent_id,
        start_date=body.start_date,
        due_date=body.due_date,
        priority=body.priority,
        status=body.status,
        progress=body.progress,
        assigned_department_id=dept_id,
        assignee_id=body.assignee_id,
        created_by_id=current_user.id,
    )
    db.add(task)
    db.flush()

    log_history(db, task.id, "created", current_user.id, {"title": task.title, "type": task.type.value})
    if task.assignee_id and task.assignee_id != current_user.id:
        assignee = db.get(User, task.assignee_id)
        log_history(db, task.id, "delegated", current_user.id, {
            "assignee_id": str(task.assignee_id),
            "assignee_name": assignee.full_name if assignee else None,
        })

    db.commit()
    db.refresh(task)

    # Best-effort Telegram notification (after commit, never raises)
    if task.assignee_id and task.assignee_id != current_user.id:
        link = task_link(task.id)
        notify_task_assignee(
            db,
            task.assignee_id,
            f'Вам назначена задача: "{task.title}"\n{link}',
        )

    return _task_out(task)


@router.get("/{task_id}")
def get_task(
    task_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    task = require_task_access(task_id, db, current_user)

    out = _task_out(task)
    out["children"] = [_task_out(c) for c in (task.children or [])]
    out["parent_chain"] = get_parent_chain(db, task)
    out["comments"] = [
        {
            "id": c.id,
            "task_id": c.task_id,
            "author_id": c.author_id,
            "author_name": c.author.full_name if c.author else None,
            "body": c.body,
            "created_at": c.created_at,
        }
        for c in (task.comments or [])[:50]
    ]
    out["history"] = [
        {
            "id": h.id,
            "task_id": h.task_id,
            "actor_id": h.actor_id,
            "actor_name": h.actor.full_name if h.actor else None,
            "event_type": h.event_type,
            "payload": h.payload,
            "at": h.at,
        }
        for h in (task.history or [])[:50]
    ]
    return out


@router.patch("/{task_id}")
def update_task(
    task_id: uuid.UUID,
    body: TaskUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    task = require_task_access(task_id, db, current_user)

    # RBAC: EMPLOYEE may only change status/progress on tasks assigned to them.
    if current_user.role == UserRole.EMPLOYEE:
        if task.assignee_id != current_user.id:
            raise HTTPException(status_code=403, detail="Можно изменять только свои задачи")
        allowed = {"status", "progress"}
        update_data = body.model_dump(exclude_unset=True)
        for key in list(update_data.keys()):
            if key not in allowed:
                del update_data[key]
    elif current_user.role == UserRole.LEAD:
        # LEAD may edit tasks within their department subtree (or their own).
        allowed_depts = visible_department_ids(db, current_user)
        has_access = (
            (allowed_depts is not None and task.assigned_department_id in allowed_depts)
            or task.assignee_id == current_user.id
            or task.created_by_id == current_user.id
        )
        if not has_access:
            raise HTTPException(status_code=403, detail="Forbidden")
        update_data = body.model_dump(exclude_unset=True)
        # …but may not move a task to another department, nor assign people
        # outside their subtree.
        if (
            "assigned_department_id" in update_data
            and update_data["assigned_department_id"] is not None
            and allowed_depts is not None
            and update_data["assigned_department_id"] not in allowed_depts
        ):
            raise HTTPException(status_code=403, detail="Можно переносить задачи только в свой отдел")
        if "assignee_id" in update_data and update_data["assignee_id"]:
            a = db.get(User, update_data["assignee_id"])
            if a and a.department_id and allowed_depts is not None and a.department_id not in allowed_depts:
                raise HTTPException(status_code=403, detail="Можно назначать только сотрудников своего отдела")
    else:
        update_data = body.model_dump(exclude_unset=True)

    old_status = task.status
    old_assignee = task.assignee_id
    old_progress = task.progress
    old_due = task.due_date
    old_dept = task.assigned_department_id

    for key, val in update_data.items():
        setattr(task, key, val)

    # Status → DONE ⇒ progress = 100
    if task.status == TaskStatus.DONE and old_status != TaskStatus.DONE:
        task.progress = 100

    # Log changes
    if "status" in update_data and update_data["status"] != old_status:
        log_history(db, task.id, "status_changed", current_user.id, {
            "old": old_status.value if old_status else None,
            "new": task.status.value,
        })

    if "assignee_id" in update_data and update_data.get("assignee_id") != old_assignee:
        new_assignee = db.get(User, task.assignee_id) if task.assignee_id else None
        log_history(db, task.id, "assignee_changed", current_user.id, {
            "old_assignee_id": str(old_assignee) if old_assignee else None,
            "new_assignee_id": str(task.assignee_id) if task.assignee_id else None,
            "new_assignee_name": new_assignee.full_name if new_assignee else None,
        })

    if "progress" in update_data and update_data["progress"] != old_progress:
        log_history(db, task.id, "progress_updated", current_user.id, {
            "old": old_progress,
            "new": task.progress,
        })

    if "due_date" in update_data and update_data.get("due_date") != old_due:
        log_history(db, task.id, "due_date_changed", current_user.id, {
            "old": old_due.isoformat() if old_due else None,
            "new": task.due_date.isoformat() if task.due_date else None,
        })

    if "assigned_department_id" in update_data and update_data.get("assigned_department_id") != old_dept:
        from app.models.department import Department
        new_dept = db.get(Department, task.assigned_department_id) if task.assigned_department_id else None
        log_history(db, task.id, "department_changed", current_user.id, {
            "old_department_id": str(old_dept) if old_dept else None,
            "new_department_id": str(task.assigned_department_id) if task.assigned_department_id else None,
            "new_department_name": new_dept.name if new_dept else None,
        })

    db.commit()
    db.refresh(task)

    # Recompute parent progress
    recompute_parent_progress(db, task)
    db.commit()

    # Best-effort Telegram notification on assignee change (after commit, never raises)
    if (
        "assignee_id" in update_data
        and task.assignee_id
        and task.assignee_id != old_assignee
        and task.assignee_id != current_user.id
    ):
        link = task_link(task.id)
        notify_task_assignee(
            db,
            task.assignee_id,
            f'Вам назначена задача: "{task.title}"\n{link}',
        )

    return _task_out(task)


@router.delete("/{task_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_task(
    task_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    task = require_task_access(task_id, db, current_user)
    # ADMIN: any task · LEAD: tasks in their department subtree · EMPLOYEE: none.
    if not can_delete_task(db, current_user, task):
        raise HTTPException(status_code=403, detail="Недостаточно прав для удаления задачи")
    db.delete(task)
    db.commit()


@router.post("/{task_id}/comments", status_code=status.HTTP_201_CREATED)
def add_comment(
    task_id: uuid.UUID,
    body: CommentCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    task = require_task_access(task_id, db, current_user)

    comment = TaskComment(
        id=uuid.uuid4(),
        task_id=task_id,
        author_id=current_user.id,
        body=body.body,
    )
    db.add(comment)
    log_history(db, task_id, "comment_added", current_user.id, {"body": body.body[:200]})
    db.commit()
    db.refresh(comment)

    # Best-effort Telegram notification to assignee (after commit, never raises)
    if task.assignee_id and task.assignee_id != current_user.id:
        link = task_link(task.id)
        notify_task_assignee(
            db,
            task.assignee_id,
            f'Новый комментарий к задаче "{task.title}":\n{link}',
        )

    return {
        "id": comment.id,
        "task_id": comment.task_id,
        "author_id": comment.author_id,
        "author_name": current_user.full_name,
        "body": comment.body,
        "created_at": comment.created_at,
    }


@router.patch("/{task_id}/comments/{comment_id}")
def update_comment(
    task_id: uuid.UUID,
    comment_id: uuid.UUID,
    body: CommentUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    require_task_access(task_id, db, current_user)
    comment = db.get(TaskComment, comment_id)
    if comment is None or comment.task_id != task_id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Комментарий не найден")
    # Only the author may edit their own comment.
    if comment.author_id != current_user.id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Можно редактировать только свои комментарии")
    comment.body = body.body
    db.commit()
    db.refresh(comment)
    return {
        "id": comment.id,
        "task_id": comment.task_id,
        "author_id": comment.author_id,
        "author_name": comment.author.full_name if comment.author else current_user.full_name,
        "body": comment.body,
        "created_at": comment.created_at,
    }


@router.delete("/{task_id}/comments/{comment_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_comment(
    task_id: uuid.UUID,
    comment_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    require_task_access(task_id, db, current_user)
    comment = db.get(TaskComment, comment_id)
    if comment is None or comment.task_id != task_id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Комментарий не найден")
    # The author may delete their own comment; ADMIN may delete any.
    if comment.author_id != current_user.id and current_user.role != UserRole.ADMIN:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Недостаточно прав для удаления комментария")
    db.delete(comment)
    db.commit()
